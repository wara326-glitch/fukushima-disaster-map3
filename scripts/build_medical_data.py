#!/usr/bin/env python3
import csv, io, json, os, re, sys, urllib.request, zipfile
from openpyxl import load_workbook
from pathlib import Path

HOSPITAL_URL = "https://www.mhlw.go.jp/content/11121000/01-1_hospital_facility_info_20260601.csv.zip"
CLINIC_URL = "https://www.mhlw.go.jp/content/11121000/02-1_clinic_facility_info_20260601.csv.zip"
CLINIC_SPECIALTY_URL = "https://www.mhlw.go.jp/content/11121000/02-2_clinic_speciality_hours_20260601.csv.zip"
BED_REPORT_URL = "https://www.mhlw.go.jp/content/10800000/001717798.xlsx"

CRITICAL = [
    "いわき市医療センター",
    "太田西ノ内病院",
    "会津中央病院",
    "福島県立医科大学附属病院",
]
DISASTER = [
    "福島県立医科大学附属病院",
    "福島赤十字病院",
    "枡記念病院",
    "太田西ノ内病院",
    "総合南東北病院",
    "公立岩瀬病院",
    "白河厚生総合病院",
    "会津中央病院",
    "福島県立南会津病院",
    "南相馬市立総合病院",
    "ふたば医療センター附属病院",
    "いわき市医療センター",
]

def norm(s):
    s = (s or "").replace("\u3000"," ").strip()
    s = re.sub(r"\s+", "", s)
    s = s.replace("―","ー").replace("－","ー").replace("ｰ","ー")
    for t in ["公立大学法人","一般財団法人","公益財団法人","医療法人","社団法人","独立行政法人","福島県厚生農業協同組合連合会","一般財団法人脳神経疾患研究所附属","一般財団法人太田綜合病院附属","一般財団法人温知会"]:
        s = s.replace(t, "")
    return s

def matches(name, targets):
    n = norm(name)
    return any(norm(t) in n for t in targets)

def download(url):
    req = urllib.request.Request(url, headers={"User-Agent":"fukushima-disaster-map3/1.0"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()

def unzip_csv(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        names = [n for n in z.namelist() if n.lower().endswith(".csv")]
        if not names:
            raise RuntimeError("CSV not found in ZIP")
        with z.open(names[0]) as f:
            return io.TextIOWrapper(f, encoding="utf-8-sig", newline="").read()

def find_col(headers, exact=(), contains=()):
    nh = {h.strip():h for h in headers}
    for e in exact:
        if e in nh:
            return nh[e]
    for h in headers:
        hs = h.strip()
        if all(x in hs for x in contains):
            return h
    return None

def first_matching(headers, patterns):
    for pat in patterns:
        for h in headers:
            if all(x in h for x in pat):
                return h
    return None

def to_num(v):
    if v is None: return None
    s = str(v).replace(",","").replace("件","").replace("台","").strip()
    if not s or s in {"-","―","－","未確認","不明"}: return None
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None

def parse_facilities(text, kind):
    reader = csv.DictReader(io.StringIO(text))
    headers = reader.fieldnames or []
    name_col = first_matching(headers, [
        ("正式名称",), ("施設名称",), ("医療機関名",), ("機関名",)
    ])
    id_col = next((h for h in headers if h.strip()=="ID"), None)
    lat_col = first_matching(headers, [("緯度",)])
    lon_col = first_matching(headers, [("経度",)])
    pref_cols = [h for h in headers if "都道府県" in h]
    city_col = first_matching(headers, [("市区町村",),("市町村",)])
    address_col = next((h for h in headers if h.strip()=="所在地"), None)
    address_cols = [address_col] if address_col else [h for h in headers if ("住所" in h or "所在地" in h) and "座標" not in h and "英語" not in h]
    phone_col = first_matching(headers, [("電話番号",),("代表電話",)])
    specialty_cols = [h for h in headers if kind=="clinic" and any(x in h for x in ["診療科","診療科目"]) and "対応" not in h]
    ambulance_cols = [h for h in headers if "救急車" in h and any(x in h for x in ["搬送","受入","受け入れ","件数","患者"])]
    if not name_col:
        raise RuntimeError("facility name column not found; first headers="+repr(headers[:30]))
    if not lat_col or not lon_col:
        raise RuntimeError("lat/lon columns not found")
    out = []
    pref_samples = []
    sample_pref07 = []
    for row in reader:
        if len(pref_samples) < 30:
            for pc in pref_cols:
                v=(row.get(pc) or "").strip()
                if v and v not in pref_samples:
                    pref_samples.append(v)
        pref_hit = any((row.get(c) or "").strip() in {"07","7","福島県"} for c in pref_cols)
        if not pref_hit:
            # fallback: search all address-like values
            pref_hit = any("福島県" in (row.get(c) or "") for c in address_cols)
        if pref_hit and len(sample_pref07) < 5:
            sample_pref07.append({
                "name": row.get(name_col),
                "lat": row.get(lat_col),
                "lon": row.get(lon_col),
                "pref": {pc: row.get(pc) for pc in pref_cols},
                "address": {ac: row.get(ac) for ac in address_cols[:8]}
            })
        if not pref_hit:
            continue
        name = (row.get(name_col) or "").strip()
        lat = to_num(row.get(lat_col)); lng = to_num(row.get(lon_col))
        if not name or lat is None or lng is None:
            continue
        address=""
        for ac in address_cols:
            if ac:
                v=(row.get(ac) or "").strip()
                if v:
                    address=v
                    break
        ambulance=None
        if kind=="hospital":
            vals=[to_num(row.get(c)) for c in ambulance_cols]
            vals=[v for v in vals if v is not None and v >= 0]
            ambulance=max(vals) if vals else None
        specialties=[]
        if kind=="clinic":
            raw=" ".join((row.get(sc) or "") for sc in specialty_cols)
            specialty_map=[
                ("内科",["内科"]),
                ("外科",["外科"]),
                ("小児科",["小児科"]),
                ("整形外科",["整形外科"]),
                ("産婦人科",["産婦人科","産科","婦人科"]),
                ("眼科",["眼科"]),
                ("耳鼻科",["耳鼻咽喉科","耳鼻科"]),
                ("精神科",["精神科","心療内科"]),
                ("皮膚科",["皮膚科"]),
                ("泌尿器科",["泌尿器科"]),
            ]
            specialties=[label for label,terms in specialty_map if any(t in raw for t in terms)]
        item={
            "name":name,
            "source_id":(row.get(id_col) or "").strip() if id_col else "",
            "lat":lat,
            "lng":lng,
            "type":kind,
            "address":address,
            "phone":(row.get(phone_col) or "").strip() if phone_col else "",
            "ambulance":int(round(ambulance)) if ambulance is not None else None,
            "specialties":specialties,
            "critical": matches(name, CRITICAL) if kind=="hospital" else False,
            "disaster": matches(name, DISASTER) if kind=="hospital" else False,
        }
        out.append(item)
    # de-duplicate by normalized name + coordinates
    dedup={}
    for x in out:
        key=(norm(x["name"]), round(x["lat"],5), round(x["lng"],5))
        dedup[key]=x
    return list(dedup.values()), {
        "name_col":name_col,"id_col":id_col,"lat_col":lat_col,"lon_col":lon_col,
        "ambulance_cols":ambulance_cols,"specialty_cols":specialty_cols,"pref_cols":pref_cols,
        "emergency_headers":[h for h in headers if "救急" in h or "搬送" in h or "救急車" in h],
        "pref_samples":pref_samples,
        "sample_pref07":sample_pref07,
        "count":len(dedup)
    }


def parse_clinic_specialties(text):
    reader=csv.DictReader(io.StringIO(text))
    headers=reader.fieldnames or []
    id_col=next((h for h in headers if h.strip()=="ID"), None)
    code_col=next((h for h in headers if h.strip()=="診療科目コード"), None)
    name_col=next((h for h in headers if h.strip()=="診療科目名"), None)
    if not id_col or not name_col:
        raise RuntimeError("specialty ID/name columns not found; headers="+repr(headers[:120]))
    specmap={}
    specialty_map=[
        ("内科",["内科"]),
        ("外科",["外科"]),
        ("小児科",["小児科"]),
        ("整形外科",["整形外科"]),
        ("産婦人科",["産婦人科","産科","婦人科"]),
        ("眼科",["眼科"]),
        ("耳鼻科",["耳鼻咽喉科","耳鼻科"]),
        ("精神科",["精神科","心療内科"]),
        ("皮膚科",["皮膚科"]),
        ("泌尿器科",["泌尿器科"]),
    ]
    for row in reader:
        sid=(row.get(id_col) or "").strip()
        sname=(row.get(name_col) or "").strip()
        if not sid or not sname:
            continue
        specs={label for label,terms in specialty_map if any(t in sname for t in terms)}
        if specs:
            specmap.setdefault(sid,set()).update(specs)
    return {k:sorted(v) for k,v in specmap.items()}, {
        "id_col":id_col,
        "code_col":code_col,
        "name_col":name_col,
        "matched_ids":len(specmap),
        "headers":headers[:40],
    }


def merge_clinic_specialties(clinics,specmap):
    matched=0
    for h in clinics:
        sid=(h.get("source_id") or "").strip()
        specs=specmap.get(sid)
        if specs:
            h["specialties"]=list(specs)
            matched+=1
        else:
            h["specialties"]=[]
    return {"matched_clinics":matched,"unmatched_clinics":len(clinics)-matched}


def parse_bed_report():
    blob = download(BED_REPORT_URL)
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [ws.cell(5,i).value for i in range(1, ws.max_column+1)]
    def col(label):
        for i,v in enumerate(headers, start=1):
            if (str(v).strip() if v is not None else "") == label:
                return i
        return None
    name_col = col("医療機関名")
    pref_col = col("都道府県コード")
    ambulance_col = col("救急車の受入件数")
    tertiary_col = col("三次救急医療施設の認定の有無")
    secondary_col = col("二次救急医療施設の認定の有無")
    if not name_col or not pref_col or not ambulance_col:
        raise RuntimeError("Required bed report columns not found")
    rows=[]
    max_needed=max(x for x in [name_col,pref_col,ambulance_col,tertiary_col,secondary_col] if x)
    for vals in ws.iter_rows(min_row=6, max_col=max_needed, values_only=True):
        pref=vals[pref_col-1]
        ps=str(pref).strip()
        if ps.endswith(".0"): ps=ps[:-2]
        if ps.zfill(2)!="07":
            continue
        name=str(vals[name_col-1] or "").strip()
        if not name:
            continue
        aval=to_num(vals[ambulance_col-1])
        rows.append({
            "name":name,
            "ambulance":int(round(aval)) if aval is not None else None,
            "tertiary":vals[tertiary_col-1] if tertiary_col else None,
            "secondary":vals[secondary_col-1] if secondary_col else None,
        })
    return rows, {
        "sheet":ws.title,
        "name_col":name_col,
        "pref_col":pref_col,
        "ambulance_col":ambulance_col,
        "tertiary_col":tertiary_col,
        "secondary_col":secondary_col,
        "fukushima_rows":len(rows),
        "with_ambulance_count":sum(1 for x in rows if x["ambulance"] is not None),
    }

def merge_ambulance_counts(hospitals, bed_rows):
    exact={norm(x["name"]):x for x in bed_rows}
    matched=0
    unmatched=[]
    for h in hospitals:
        hn=norm(h["name"])
        hit=exact.get(hn)
        if not hit:
            candidates=[]
            for bn,row in exact.items():
                if min(len(hn),len(bn)) >= 6 and (hn in bn or bn in hn):
                    candidates.append((abs(len(hn)-len(bn)),row))
            if candidates:
                candidates.sort(key=lambda z:z[0])
                hit=candidates[0][1]
        if hit:
            h["ambulance"]=hit.get("ambulance")
            h["bed_report_name"]=hit.get("name")
            matched+=1
        else:
            unmatched.append(h["name"])
    return {"matched_hospitals":matched,"unmatched_hospitals":unmatched}

def category(x):
    if x["type"]=="clinic": return "clinic"
    if x.get("critical"): return "critical"
    if x.get("disaster"): return "disaster"
    a=x.get("ambulance")
    if a is not None and a>=1000: return "ambulance1000"
    if a is not None and a>=500: return "ambulance500"
    return "hospital"

def main():
    hospitals, hm = parse_facilities(unzip_csv(download(HOSPITAL_URL)), "hospital")
    clinics, cm = parse_facilities(unzip_csv(download(CLINIC_URL)), "clinic")
    specmap, spec_meta = parse_clinic_specialties(unzip_csv(download(CLINIC_SPECIALTY_URL)))
    spec_merge = merge_clinic_specialties(clinics, specmap)
    bed_rows, bed_meta=parse_bed_report()
    merge_meta=merge_ambulance_counts(hospitals, bed_rows)
    facilities=hospitals+clinics
    for x in facilities:
        x["category"]=category(x)
    counts={}
    for x in facilities:
        counts[x["category"]]=counts.get(x["category"],0)+1
    payload={
        "source":"厚生労働省 医療情報ネット オープンデータ",
        "as_of":"2026-06-01",
        "generated_from":{"hospital":HOSPITAL_URL,"clinic":CLINIC_URL,"clinic_specialty":CLINIC_SPECIALTY_URL,"bed_report":BED_REPORT_URL},
        "ambulance_period":"2024-04-01/2025-03-31",
        "counts":counts,
        "facilities":facilities,
    }
    Path("data").mkdir(exist_ok=True)
    Path("data/fukushima_medical.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",",":")),encoding="utf-8")
    ambulance_rank=sorted(
        [{"name":h["name"],"ambulance":h.get("ambulance"),"category":h.get("category")} for h in hospitals if h.get("ambulance") is not None],
        key=lambda x:x["ambulance"], reverse=True
    )
    meta={"hospital":hm,"clinic":cm,"clinic_specialty":spec_meta,"clinic_specialty_merge":spec_merge,"bed_report":bed_meta,"merge":merge_meta,"counts":counts,"ambulance_rank":ambulance_rank}
    Path("data/build-meta.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False, indent=2))

if __name__=="__main__":
    main()
